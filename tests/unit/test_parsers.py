import io
from datetime import date, datetime
from typing import Any

import docx
import openpyxl
import pytest
from docx.oxml import OxmlElement

from app.core.errors import UnreadableDocumentError
from app.parsers.blocks import build_blocks
from app.parsers.pdf import NO_TEXT_LAYER_REASON
from app.parsers.xlsx import format_cell
from app.schemas.documents import (
    Chunk,
    DocumentFormat,
    HeadedLine,
    TableBlock,
    TableRow,
    TextBlock,
)
from app.services.documents import read_chunks
from tests.fixtures import DOCUMENTS_DIR

SCAN = "fan_passport_fan_45_scan.pdf"
READABLE_FIXTURES = sorted(
    path.name for path in DOCUMENTS_DIR.iterdir() if path.name != SCAN
)
FAN_40_ROW = (
    "FAN-40: Brand VentBriz; Model 40; Power 40 W; Blade diameter 27 cm; "
    "Weight 2.5 kg; Wholesale price 2,490 RUB"
)


def read_fixture(name: str) -> list[Chunk]:
    """Return the chunks of a fixture document"""
    path = DOCUMENTS_DIR / name
    document_format = DocumentFormat(path.suffix.removeprefix("."))
    chunks, _ = read_chunks(document_format, path.read_bytes())
    return chunks


def find_chunk(chunks: list[Chunk], text: str) -> Chunk:
    """Return the first chunk that contains the text"""
    return next(chunk for chunk in chunks if text in chunk.text)


def save_docx(document: Any) -> bytes:
    """Return the python-docx document as file bytes"""
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_reads_sections_page_by_page() -> None:
    """Chunks follow the passport's headings across its two pages"""
    chunks = read_fixture("fan_passport_fan_40.pdf")

    placement = [
        (chunk.metadata["page"], chunk.metadata["section"])
        for chunk in chunks[:5]
    ]
    power = find_chunk(chunks, "Power 40 W")

    assert placement == [
        (1, "VentBriz 40 Fan"),
        (1, "Technical specifications"),
        (2, "Package contents"),
        (2, "Operating rules"),
        (2, "Warranty obligations"),
    ]
    assert power.metadata["page"] == 1
    assert power.metadata["section"] == "Technical specifications"


def test_pdf_drops_running_header() -> None:
    """The company line repeated on both pages reaches no chunk"""
    chunks = read_fixture("fan_passport_fan_40.pdf")

    assert not any("opt@domteh.example" in chunk.text for chunk in chunks)


def test_pdf_drops_numbered_footer() -> None:
    """The footer differs only by page number and is still removed"""
    chunks = read_fixture("coffee_passport_cfe_1000.pdf")

    placement = {
        (chunk.metadata["page"], chunk.metadata["section"]) for chunk in chunks
    }

    assert not any("sales@electromir.example" in chunk.text for chunk in chunks)
    assert (1, "Технические характеристики") in placement


def test_single_page_pdf_keeps_company_line() -> None:
    """Nothing repeats in a one-page document, so its company line stays"""
    chunks = read_fixture("fan_kp_2.pdf")

    assert find_chunk(chunks, "DomTekh LLC · opt@domteh.example")


def test_xlsx_row_names_every_column() -> None:
    """Each spreadsheet row is one chunk that carries its column headers"""
    chunks = read_fixture("fan_spec_1.xlsx")

    fan_40 = find_chunk(chunks, "FAN-40")

    assert [chunk.metadata["kind"] for chunk in chunks] == ["table_row"] * 3
    assert fan_40.text == FAN_40_ROW
    assert fan_40.metadata == {
        "position": 1,
        "kind": "table_row",
        "page": 1,
        "section": "Specification",
        "row": 3,
        "sku": "FAN-40",
        "brand": "VentBriz",
    }


def test_xlsx_keeps_documented_power_apart_from_sku() -> None:
    """The number in VCS-180 is not its power"""
    chunks = read_fixture("vacuum_spec_1.xlsx")

    assert "Power 500 W" in find_chunk(chunks, "VCS-180").text


@pytest.mark.parametrize(
    ("name", "section"),
    [
        ("blender_kp.docx", "Price list (wholesale)"),
        ("blender_kp_ru.docx", "Прайс-лист (опт)"),
    ],
)
def test_docx_table_rows_carry_skus(name: str, section: str) -> None:
    """Price-list rows are table rows with their SKU under their heading"""
    chunks = read_fixture(name)

    rows = [chunk for chunk in chunks if chunk.metadata["kind"] == "table_row"]

    assert [row.metadata.get("sku") for row in rows] == [
        "BLD-500",
        "BLD-800",
        "BLD-1200",
    ]
    assert {row.metadata["section"] for row in rows} == {section}
    assert [row.metadata.get("row") for row in rows] == [2, 3, 4]
    # The heading right above the table names the rows' section and makes no
    # chunk of its own.
    assert all(chunk.text != section for chunk in chunks)


def test_scan_without_text_layer_is_unreadable() -> None:
    """An image-only PDF fails with the no-text-layer reason"""
    with pytest.raises(UnreadableDocumentError, match=NO_TEXT_LAYER_REASON):
        read_fixture(SCAN)


@pytest.mark.parametrize("name", READABLE_FIXTURES)
def test_every_chunk_has_position_page_and_section(name: str) -> None:
    """Chunks are numbered without gaps, and each has a page and section"""
    chunks = read_fixture(name)

    assert [chunk.metadata["position"] for chunk in chunks] == list(
        range(len(chunks))
    )
    assert all(chunk.metadata["page"] >= 1 for chunk in chunks)
    assert all(chunk.metadata["section"] for chunk in chunks)
    for chunk in chunks:
        if chunk.metadata["kind"] == "table_row":
            assert chunk.text.startswith(f"{chunk.metadata.get('sku')}: ")


def test_consecutive_headings_name_one_section() -> None:
    """The last heading of a group names the section; headings stay"""
    blocks = build_blocks(
        [
            HeadedLine(1, "PRODUCT PASSPORT", True),
            HeadedLine(1, "VentBriz 40 Fan", True),
            HeadedLine(1, "Floor fan.", False),
        ]
    )

    assert blocks == [
        TextBlock(
            page=1,
            section="VentBriz 40 Fan",
            lines=["PRODUCT PASSPORT", "VentBriz 40 Fan", "Floor fan."],
        )
    ]


def test_sections_continue_across_pages_and_tables() -> None:
    """Text before the first heading joins it; pages and tables keep it"""
    table = TableBlock(
        page=2,
        section=None,
        headers=["SKU"],
        rows=[TableRow(number=2, cells=["FAN-40"])],
    )

    blocks = build_blocks(
        [
            HeadedLine(1, "DomTekh LLC", False),
            HeadedLine(1, "Specifications", True),
            HeadedLine(1, "Power 40 W", False),
            HeadedLine(2, "Weight 2.5 kg", False),
            table,
        ]
    )

    assert [(block.page, block.section) for block in blocks] == [
        (1, "Specifications"),
        (1, "Specifications"),
        (2, "Specifications"),
        (2, "Specifications"),
    ]


def test_headings_alone_make_no_block() -> None:
    """A heading right above a table, or at the end, adds no text block"""
    table = TableBlock(
        page=1,
        section=None,
        headers=["SKU"],
        rows=[TableRow(number=2, cells=["BLD-800"])],
    )

    blocks = build_blocks(
        [
            HeadedLine(1, "Price list", True),
            table,
            HeadedLine(1, "Terms", True),
        ]
    )

    assert blocks == [table.model_copy(update={"section": "Price list"})]


def test_document_without_headings_has_no_section() -> None:
    """A document without any heading has a null section"""
    blocks = build_blocks([HeadedLine(1, "Power 40 W", False)])

    assert [block.section for block in blocks] == [None]


def test_docx_page_follows_explicit_page_breaks() -> None:
    """Without layout marks, explicit page breaks set the page"""
    document = docx.Document()
    document.add_heading("Warranty", level=1)
    document.add_paragraph("Twelve months.")
    document.add_page_break()
    document.add_paragraph("Service centres.")

    chunks, page_count = read_chunks(DocumentFormat.DOCX, save_docx(document))

    assert [
        (chunk.text, chunk.metadata["page"], chunk.metadata["section"])
        for chunk in chunks
    ] == [
        ("Warranty\nTwelve months.", 1, "Warranty"),
        ("Service centres.", 2, "Warranty"),
    ]
    assert page_count == 2


def test_docx_page_follows_only_layout_marks_when_present() -> None:
    """Word's layout marks win over explicit page breaks"""
    document = docx.Document()
    document.add_paragraph("One")
    # Word marks a page that starts with a paragraph inside it, before its
    # text.
    starts_page = document.add_paragraph("Two")
    starts_page.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
    document.add_page_break()
    document.add_paragraph("Three")

    chunks, _ = read_chunks(DocumentFormat.DOCX, save_docx(document))

    assert [(chunk.text, chunk.metadata["page"]) for chunk in chunks] == [
        ("One", 1),
        ("Two\nThree", 2),
    ]


def test_xlsx_reads_visible_sheets_at_their_position() -> None:
    """Hidden sheets are skipped, and empty leading columns are not read"""
    workbook = openpyxl.Workbook()
    hidden = workbook.active
    assert hidden is not None
    hidden.title = "Hidden"
    hidden["A1"] = "secret"
    hidden.sheet_state = "hidden"
    sheet = workbook.create_sheet("Prices")
    sheet.append([])
    sheet.append([None, "Model", "Weight, kg", "Since", "Total"])
    sheet.append([])
    sheet.append([None, "VentBriz 35", 2.5, date(2026, 1, 2), "=C4*2"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    chunks, page_count = read_chunks(DocumentFormat.XLSX, buffer.getvalue())

    assert [(chunk.text, chunk.metadata) for chunk in chunks] == [
        (
            "VentBriz 35: Weight 2.5 kg; Since 2026-01-02",
            {
                "position": 0,
                "kind": "table_row",
                "page": 2,
                "section": "Prices",
                "row": 4,
            },
        )
    ]
    assert page_count == 1


@pytest.mark.parametrize(
    ("value", "text"),
    [
        (35.0, "35"),
        (35, "35"),
        (2.4, "2.4"),
        # Excel keeps no time zone, so openpyxl reads naive datetimes.
        (datetime.fromisoformat("2026-01-02T00:00"), "2026-01-02"),
        (datetime.fromisoformat("2026-01-02T09:30"), "2026-01-02T09:30:00"),
        (date(2026, 1, 2), "2026-01-02"),
        ("  black ", "black"),
        (None, ""),
    ],
)
def test_format_cell(value: object, text: str) -> None:
    """Cell values read the way a person sees them"""
    assert format_cell(value) == text
