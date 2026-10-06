import io
from datetime import date, datetime, time
from typing import Any

import openpyxl

from app.core.errors import UnreadableDocumentError
from app.schemas.documents import (
    ParsedBlock,
    ParsedDocument,
    TableBlock,
    TableRow,
)

UNREADABLE_XLSX_REASON = "the file is not a readable XLSX"
VISIBLE_SHEET_STATE = "visible"
FIRST_ROW_NUMBER = 1
FIRST_SHEET_POSITION = 1

SheetValues = tuple[int, str, list[tuple[Any, ...]]]


def format_cell(value: Any) -> str:
    """Return the cell's value as text"""
    if value is None:
        return ""
    # Excel stores numbers as floats, and "35" must not read "35.0". str()
    # gives the shortest form of any other number: 2.4, not 2.39999...
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, datetime):
        is_date_only = value.time() == time()
        return value.date().isoformat() if is_date_only else value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value).strip()


def read_sheets(content: bytes) -> list[SheetValues]:
    """Return the position, name and cell values of each visible sheet"""
    # Read-only mode streams rows instead of loading every cell's styles.
    # data_only gives a formula's last saved value, or None without one.
    workbook = openpyxl.load_workbook(
        io.BytesIO(content), read_only=True, data_only=True
    )
    try:
        sheets = [
            (position, sheet.title, list(sheet.iter_rows(values_only=True)))
            for position, sheet in enumerate(
                workbook.worksheets, start=FIRST_SHEET_POSITION
            )
            if sheet.sheet_state == VISIBLE_SHEET_STATE
        ]
    finally:
        workbook.close()
    return sheets


def read_table(
    position: int, name: str, values: list[tuple[Any, ...]]
) -> TableBlock | None:
    """Return the sheet's table, or None if every cell is empty"""
    # Read-only rows start at sheet row 1, empty rows included, so the
    # position in the list gives the row number a person sees.
    filled: list[tuple[int, list[str]]] = []
    for number, row in enumerate(values, start=FIRST_ROW_NUMBER):
        cells = [format_cell(value) for value in row]
        if any(cells):
            filled.append((number, cells))
    if not filled:
        return None
    # A table often starts right of column A. Columns empty in every row are
    # not part of it, so "the first column" means the table's first.
    first_column = min(
        next(index for index, cell in enumerate(cells) if cell)
        for _, cells in filled
    )
    (_, headers), *rows = [
        (number, cells[first_column:]) for number, cells in filled
    ]
    return TableBlock(
        page=position,
        section=name,
        headers=headers,
        rows=[TableRow(number=number, cells=cells) for number, cells in rows],
    )


def parse_xlsx(content: bytes) -> ParsedDocument:
    """Return one table block per visible sheet that has a value"""
    try:
        sheets = read_sheets(content)
    # openpyxl raises many error types for a broken file. Any of them means
    # the file can't be read, and retrying won't change that.
    except Exception as error:
        raise UnreadableDocumentError(UNREADABLE_XLSX_REASON) from error
    blocks: list[ParsedBlock] = []
    for position, name, values in sheets:
        table = read_table(position, name, values)
        if table is not None:
            blocks.append(table)
    # Hidden sheets are not read, so they don't count.
    return ParsedDocument(blocks=blocks, page_count=len(sheets))
