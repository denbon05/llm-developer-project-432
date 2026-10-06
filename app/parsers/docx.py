import io

import docx
from docx.document import Document
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.table import Table

from app.core.errors import UnreadableDocumentError
from app.parsers.blocks import build_blocks
from app.schemas.documents import (
    HeadedLine,
    ParsedDocument,
    TableBlock,
    TableRow,
)

UNREADABLE_DOCX_REASON = "the file is not a readable DOCX"
# python-docx reports built-in style names in English, whatever the language
# of the Word installation that saved the file.
HEADING_STYLES = frozenset(
    {"Title", *(f"Heading {level}" for level in range(1, 10))}
)
FIRST_PAGE = 1
HEADER_ROW_NUMBER = 1
CELL_PARAGRAPH_SEPARATOR = " "
TEXT_TAG = qn("w:t")
BREAK_TAG = qn("w:br")
BREAK_TYPE_ATTRIBUTE = qn("w:type")
PAGE_BREAK_TYPE = "page"
RENDERED_PAGE_BREAK_TAG = qn("w:lastRenderedPageBreak")


def is_page_break(node: BaseOxmlElement, is_rendered: bool) -> bool:
    """Return whether the node is a page break of the kind that counts"""
    if is_rendered:
        return node.tag == RENDERED_PAGE_BREAK_TAG
    return (
        node.tag == BREAK_TAG
        and node.get(BREAK_TYPE_ATTRIBUTE) == PAGE_BREAK_TYPE
    )


def count_page_breaks(
    element: BaseOxmlElement, is_rendered: bool
) -> tuple[int, int]:
    """Return the page breaks before the element's first text, and all"""
    # Word puts the mark of a page that starts with a paragraph inside that
    # paragraph, before its text, so those breaks move the paragraph itself.
    leading = total = 0
    has_text = False
    for node in element.iter(TEXT_TAG, BREAK_TAG, RENDERED_PAGE_BREAK_TAG):
        if node.tag == TEXT_TAG:
            has_text = has_text or bool(node.text)
        elif is_page_break(node, is_rendered):
            total += 1
            if not has_text:
                leading += 1
    return leading, total


def read_table(table: Table, page: int) -> TableBlock:
    """Return the table with its first row as the column headers"""
    # A merged cell repeats under every column it spans: python-docx returns
    # the same cell for each of them.
    rows = [
        [
            CELL_PARAGRAPH_SEPARATOR.join(
                paragraph.text for paragraph in cell.paragraphs
            )
            for cell in row.cells
        ]
        for row in table.rows
    ]
    return TableBlock(
        page=page,
        section=None,
        headers=rows[0] if rows else [],
        rows=[
            TableRow(number=number, cells=cells)
            for number, cells in enumerate(
                rows[1:], start=HEADER_ROW_NUMBER + 1
            )
        ],
    )


def read_items(document: Document) -> tuple[list[HeadedLine | TableBlock], int]:
    """Return the body's paragraphs and tables in order, and the page count"""
    # DOCX stores no page numbers. When Word last laid out the file, it marked
    # where each page began; without those marks, explicit page breaks are
    # the only hint. A generated file with neither is page 1 throughout.
    is_rendered = any(
        True for _ in document.element.body.iter(RENDERED_PAGE_BREAK_TAG)
    )
    items: list[HeadedLine | TableBlock] = []
    breaks_before = 0
    # Page headers and footers are separate parts of the file, not the body.
    for item in document.iter_inner_content():
        leading, total = count_page_breaks(item._element, is_rendered)
        page = FIRST_PAGE + breaks_before + leading
        breaks_before += total
        if isinstance(item, Table):
            items.append(read_table(item, page))
        # A blank paragraph would split a group of consecutive headings.
        elif item.text.strip():
            is_heading = (
                item.style is not None and item.style.name in HEADING_STYLES
            )
            items.append(HeadedLine(page, item.text, is_heading))
    return items, FIRST_PAGE + breaks_before


def parse_docx(content: bytes) -> ParsedDocument:
    """Return the text and table blocks of a DOCX body, in document order"""
    try:
        items, page_count = read_items(docx.Document(io.BytesIO(content)))
    # python-docx raises many error types for a broken file. Any of them
    # means the file can't be read, and retrying won't change that.
    except Exception as error:
        raise UnreadableDocumentError(UNREADABLE_DOCX_REASON) from error
    return ParsedDocument(blocks=build_blocks(items), page_count=page_count)
